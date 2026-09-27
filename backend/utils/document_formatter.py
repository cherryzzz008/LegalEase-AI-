from io import BytesIO

from docx import Document
from fpdf import FPDF


def format_txt(text, document_type):
    return f"{document_type}\n\n{text}".encode("utf-8")


def format_docx(text, document_type, logo_path=None, terms=None):
    document = Document()

    if logo_path:
        try:
            document.add_picture(logo_path, width=None)
        except Exception:
            pass

    document.add_heading(document_type, level=1)

    for paragraph in text.split("\n"):
        document.add_paragraph(paragraph)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output.getvalue()


def format_pdf(text, document_type, logo_path=None):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, document_type)
    pdf.ln(5)

    pdf.set_font("Helvetica", size=11)

    for paragraph in text.split("\n"):
        pdf.multi_cell(0, 7, paragraph)
        pdf.ln(2)

    return bytes(pdf.output())